from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"C:\nakiuebe\firstjob\roxlease\docs\ocr\office-lease-blank.docx"
BLANK = "........................................................"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.1)
sec.left_margin = Cm(2.7)
sec.right_margin = Cm(2.2)
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(5)

for name, size, before, after in [("Article", 12, 14, 7), ("Annex", 13, 0, 12)]:
    st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = normal
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


def p(text="", *, bold_lead=None, align=None, keep=False, before=0, after=None):
    para = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    if align is not None:
        para.alignment = align
    para.paragraph_format.keep_together = keep
    para.paragraph_format.space_before = Pt(before)
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    return para


def field(label, space=BLANK):
    para = doc.add_paragraph()
    para.add_run(label + ": ").bold = True
    para.add_run(space)
    para.paragraph_format.left_indent = Cm(0.45)
    para.paragraph_format.space_after = Pt(3)
    return para


def article(n, title):
    para = doc.add_paragraph(style="Article")
    para.add_run(f"ĐIỀU {n}. {title}")


def clause(n, text):
    return p(f"{n}. {text}", keep=True)


def annex(n, title):
    doc.add_page_break()
    para = doc.add_paragraph(style="Annex")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run(f"PHỤ LỤC {n}\n{title}")
    p("Kèm theo Hợp đồng số " + BLANK + " ký ngày " + BLANK, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)


def subsection(n, title):
    para = doc.add_paragraph()
    para.add_run(f"{n}. {title}").bold = True
    para.paragraph_format.space_before = Pt(9)
    para.paragraph_format.keep_with_next = True


def signatures():
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(20)
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.tab_stops.add_tab_stop(Cm(8.3))
    para.add_run("ĐẠI DIỆN BÊN A\tĐẠI DIỆN BÊN B").bold = True
    para = doc.add_paragraph()
    para.paragraph_format.tab_stops.add_tab_stop(Cm(8.3))
    para.add_run("(Ký, ghi rõ họ tên, đóng dấu nếu có)\t(Ký, ghi rõ họ tên, đóng dấu nếu có)").italic = True
    para.paragraph_format.space_after = Pt(35)


p("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", align=WD_ALIGN_PARAGRAPH.CENTER, after=0).runs[0].bold = True
p("Độc lập - Tự do - Hạnh phúc", align=WD_ALIGN_PARAGRAPH.CENTER, after=20).runs[0].bold = True
t = p("HỢP ĐỒNG THUÊ VĂN PHÒNG", align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
t.runs[0].bold = True
t.runs[0].font.size = Pt(16)
p("Số: " + BLANK, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
p("Căn cứ quy định pháp luật Việt Nam hiện hành và nhu cầu, thỏa thuận của các bên;")
p("Hôm nay, ngày ........ tháng ........ năm ........, tại " + BLANK + ", chúng tôi gồm:")

subsection("A", "BÊN CHO THUÊ (BÊN A)")
for x in ["Tên tổ chức/cá nhân", "Mã số thuế (nếu có)", "Địa chỉ đăng ký", "Địa chỉ liên hệ", "Người đại diện", "Chức vụ", "Giấy tờ định danh", "Điện thoại", "Email", "Số tài khoản", "Ngân hàng"]:
    field(x)

subsection("B", "BÊN THUÊ (BÊN B)")
for x in ["Tên tổ chức/cá nhân", "Mã số thuế (nếu có)", "Địa chỉ đăng ký", "Địa chỉ liên hệ", "Người đại diện", "Chức vụ", "Giấy tờ định danh", "Điện thoại", "Email", "Số tài khoản", "Ngân hàng"]:
    field(x)
p("Sau khi trao đổi, hai bên thống nhất ký Hợp đồng thuê văn phòng với các điều khoản sau:", before=10)

article(1, "ĐỐI TƯỢNG VÀ MỤC ĐÍCH THUÊ")
clause("1.1", "Bên A cho Bên B thuê diện tích văn phòng tại địa chỉ " + BLANK + ". Tên dự án/khu nhà: " + BLANK + "; tòa nhà: " + BLANK + "; tầng: " + BLANK + "; phòng/suite: " + BLANK + ".")
clause("1.2", "Diện tích thuê: ........ m²; diện tích sử dụng: ........ m²; diện tích hành lang phân bổ: ........ m²; diện tích thương lượng/tính giá: ........ m². Ranh giới và cách đo diện tích theo bản vẽ/tài liệu số " + BLANK + " tại Phụ lục 1.")
clause("1.3", "Mục đích sử dụng: " + BLANK + ". Bên B dùng diện tích thuê đúng mục đích, tuân thủ quy định tòa nhà và pháp luật áp dụng.")
clause("1.4", "Hình thức thuê chính/cho thuê lại: " + BLANK + ". Nếu cho thuê lại, hợp đồng gốc hoặc căn cứ quyền cho thuê số: " + BLANK + ". Danh sách từng mặt bằng và tiện ích ghi tại Phụ lục 1.")

article(2, "THỜI HẠN VÀ BÀN GIAO")
clause("2.1", "Ngày ký: ........ / ........ / ........; ngày bàn giao dự kiến: ........ / ........ / ........; ngày bắt đầu thuê: ........ / ........ / ........; ngày kết thúc thuê: ........ / ........ / .........")
clause("2.2", "Thời gian miễn tiền thuê (nếu có): " + BLANK + ". Điều kiện bắt đầu tính tiền thuê: " + BLANK + ".")
clause("2.3", "Bên A bàn giao khu vực thuê theo biên bản tại Phụ lục 3, ghi tình trạng, chỉ số công tơ, thiết bị, chìa khóa và thẻ. Thay đổi ngày bàn giao hoặc thời hạn thuê phải được hai bên xác nhận bằng văn bản.")

article(3, "GIÁ THUÊ, THUẾ VÀ CÁC KHOẢN PHÍ")
clause("3.1", "Đồng tiền thanh toán: " + BLANK + "; loại giá thuê/đơn vị tính: " + BLANK + "; đơn giá thuê: " + BLANK + " /m²/kỳ; đơn giá dịch vụ: " + BLANK + " /m²/kỳ; giá thuê kỳ đầu: " + BLANK + ".")
clause("3.2", "Giá nêu trên đã bao gồm/chưa bao gồm thuế GTGT: " + BLANK + "; thuế suất áp dụng: " + BLANK + ". Khi thanh toán bằng đồng tiền khác, tỷ giá và nguồn tỷ giá áp dụng theo Phụ lục 2.")
clause("3.3", "Các khoản phí định kỳ, điện, nước, viễn thông, gửi xe và dịch vụ phát sinh được xác định tại Phụ lục 2 hoặc theo chứng từ và đơn giá áp dụng thực tế. Mọi điều chỉnh giá/phí phải ghi rõ thời điểm hiệu lực và được hai bên xác nhận bằng văn bản.")

article(4, "ĐẶT CỌC VÀ THANH TOÁN")
clause("4.1", "Tiền đặt cọc: " + BLANK + "; hạn nộp: ........ / ........ / ........; số kỳ thanh toán trước: " + BLANK + "; kỳ thanh toán: " + BLANK + "; hạn thanh toán mỗi kỳ: " + BLANK + ".")
clause("4.2", "Bên B thanh toán bằng phương thức " + BLANK + " vào tài khoản thụ hưởng " + BLANK + " tại ngân hàng " + BLANK + ". Bên A cung cấp chứng từ/hóa đơn theo quy định áp dụng.")
clause("4.3", "Điều kiện khấu trừ tiền cọc, thời hạn hoàn cọc và cách xử lý chậm thanh toán được ghi tại Phụ lục 4. Tiền cọc bảo đảm nghĩa vụ hợp đồng, không mặc nhiên được dùng thay tiền thuê.")

article(5, "SỬ DỤNG, BẢO TRÌ VÀ CẢI TẠO")
clause("5.1", "Bên B sử dụng tài sản đúng mục đích, giữ vệ sinh, an toàn và phòng cháy chữa cháy; không thay đổi kết cấu khi chưa được Bên A chấp thuận bằng văn bản.")
clause("5.2", "Tình trạng mặt bằng, tài sản bàn giao và trách nhiệm bảo trì/sửa chữa từng hạng mục ghi tại Phụ lục 3.")
clause("5.3", "Việc cải tạo, lắp đặt hoặc tháo dỡ thiết bị phải tuân thủ điều kiện chấp thuận và cách bàn giao lại nêu tại Phụ lục 4.")

article(6, "QUYỀN VÀ NGHĨA VỤ CỦA BÊN A")
clause("6.1", "Bàn giao đúng khu vực thuê, thời điểm và tình trạng đã thỏa thuận; bảo đảm quyền cho thuê và việc sử dụng ổn định của Bên B trong thời hạn hợp đồng.")
clause("6.2", "Bảo trì, sửa chữa phần thuộc trách nhiệm Bên A; thông báo hợp lý khi cần vào khu vực thuê, trừ tình huống khẩn cấp.")
clause("6.3", "Nhận thanh toán đúng hạn; yêu cầu Bên B khắc phục vi phạm theo trình tự thông báo và thời hạn đã thỏa thuận.")

article(7, "QUYỀN VÀ NGHĨA VỤ CỦA BÊN B")
clause("7.1", "Nhận và sử dụng khu vực thuê, tiện ích theo hợp đồng; yêu cầu Bên A khắc phục hư hỏng thuộc trách nhiệm Bên A.")
clause("7.2", "Thanh toán đúng hạn, bảo quản tài sản, tuân thủ quy định tòa nhà; chịu trách nhiệm về thiệt hại do mình gây ra theo thỏa thuận và pháp luật áp dụng.")
clause("7.3", "Chỉ cho thuê lại, chuyển giao quyền thuê hoặc thay đổi mục đích sử dụng khi đáp ứng điều kiện chấp thuận bằng văn bản tại Phụ lục 4.")

article(8, "GIA HẠN, MỞ RỘNG VÀ QUYỀN LỰA CHỌN")
clause("8.1", "Quyền gia hạn, mở rộng, chấm dứt sớm hoặc lựa chọn khác chỉ phát sinh khi được ghi cụ thể tại Phụ lục 4, gồm bên thực hiện, điều kiện, thời hạn thông báo và chi phí (nếu có).")
clause("8.2", "Khi thực hiện quyền lựa chọn, mọi thay đổi về diện tích, thời hạn hoặc giá được hai bên xác nhận bằng văn bản/phụ lục.")

article(9, "VI PHẠM, CHẤM DỨT VÀ BẤT KHẢ KHÁNG")
clause("9.1", "Căn cứ chấm dứt trước hạn, thời hạn báo trước, thời gian khắc phục vi phạm và mức bồi thường/phạt (nếu có) ghi tại Phụ lục 4.")
clause("9.2", "Bên bị ảnh hưởng bởi sự kiện bất khả kháng thông báo cho bên kia, phối hợp hạn chế thiệt hại và thống nhất xử lý nghĩa vụ bị ảnh hưởng theo quy định pháp luật và thỏa thuận cụ thể.")
clause("9.3", "Khi kết thúc hợp đồng, hai bên đối chiếu công nợ, bàn giao lại tài sản và xử lý tiền cọc theo biên bản thanh lý.")

article(10, "THÔNG BÁO, TRANH CHẤP VÀ HIỆU LỰC")
clause("10.1", "Địa chỉ, đầu mối và phương thức nhận thông báo của mỗi bên ghi tại thông tin các bên và Phụ lục 5. Thay đổi đầu mối phải được báo bằng văn bản.")
clause("10.2", "Tranh chấp được ưu tiên giải quyết bằng thương lượng. Nếu không giải quyết được, cơ quan giải quyết tranh chấp và luật áp dụng: " + BLANK + ".")
clause("10.3", "Hợp đồng có hiệu lực từ ngày " + BLANK + " hoặc theo điều kiện " + BLANK + "; được lập thành ........ bản có giá trị như nhau, mỗi bên giữ ........ bản.")
clause("10.4", "Phụ lục, biên bản bàn giao và văn bản sửa đổi được ký hợp lệ là bộ phận của hợp đồng. Thứ tự ưu tiên khi nội dung khác nhau: " + BLANK + ".")
signatures()

annex(1, "MẶT BẰNG VÀ TIỆN ÍCH")
subsection("1", "Thông tin mặt bằng")
for x in ["Địa chỉ tài sản", "Tên dự án/khu nhà", "Tòa nhà", "Tầng", "Phòng/suite", "Diện tích thuê (m²)", "Diện tích sử dụng (m²)", "Diện tích hành lang phân bổ (m²)", "Diện tích thương lượng/tính giá (m²)", "Ngày bắt đầu sử dụng", "Ngày kết thúc sử dụng", "Mục đích sử dụng", "Số hiệu bản vẽ/biên bản kèm theo"]:
    field(x)
p("Nếu hợp đồng gồm nhiều khu vực thuê, hai bên mô tả tiếp từng khu vực bằng cùng các thông tin trên và ký xác nhận vào trang bổ sung.", before=5)
subsection("2", "Tiện ích và dịch vụ đi kèm")
for x in ["Tên tiện ích/dịch vụ", "Mã hoặc hạng mục", "Thời gian áp dụng", "Điều kiện sử dụng/ghi chú"]:
    field(x)
p("Các tiện ích khác (nếu có): " + BLANK)
signatures()

annex(2, "GIÁ THUÊ VÀ CHI PHÍ ĐỊNH KỲ")
subsection("1", "Giá thuê")
for x in ["Đồng tiền thanh toán", "Loại giá thuê và đơn vị tính", "Đơn giá thuê/m²/kỳ", "Đơn giá dịch vụ/m²/kỳ", "Diện tích tính tiền thuê", "Giá thuê kỳ đầu", "Thuế GTGT đã gồm/chưa gồm", "Thuế suất áp dụng", "Ngày bắt đầu tính tiền thuê", "Thời điểm và cách điều chỉnh giá", "Tỷ giá áp dụng/nguồn tỷ giá (nếu có)"]:
    field(x)
subsection("2", "Khoản thu định kỳ")
for x in ["Tên khoản thu", "Đơn giá hoặc số tiền", "Thuế GTGT", "Chu kỳ thanh toán", "Ngày bắt đầu", "Ngày kết thúc", "Bên chịu phí"]:
    field(x)
p("Nếu có nhiều khoản thu, hai bên ghi tiếp từng khoản theo các mục trên và ký xác nhận.")
subsection("3", "Chi phí sử dụng thực tế")
field("Cách tính và bên chịu tiền điện, nước, internet, gửi xe, dịch vụ khác")
field("Chứng từ hoặc đơn giá dùng để đối chiếu")
signatures()

annex(3, "BIÊN BẢN BÀN GIAO")
for x in ["Ngày, giờ bàn giao", "Địa điểm bàn giao", "Người giao", "Người nhận", "Tình trạng chung của mặt bằng", "Chỉ số công tơ điện", "Chỉ số đồng hồ nước", "Số lượng chìa khóa", "Số lượng thẻ ra vào", "Danh mục thiết bị/tài sản bàn giao", "Số lượng và tình trạng từng hạng mục", "Bên chịu trách nhiệm bảo trì/sửa chữa", "Ảnh hoặc tài liệu kèm theo", "Các tồn tại cần khắc phục và thời hạn"]:
    field(x)
p("Hai bên xác nhận đã kiểm tra hiện trạng, số lượng và chứng từ kèm theo tại thời điểm bàn giao.", before=10)
signatures()

annex(4, "ĐIỀU KHOẢN VÀ QUYỀN LỰA CHỌN BỔ SUNG")
for x in ["Thời hạn báo gia hạn", "Điều kiện và cách điều chỉnh giá khi gia hạn", "Điều kiện mở rộng hoặc thu hẹp diện tích", "Điều kiện cho thuê lại/chuyển giao quyền thuê", "Điều kiện cải tạo, lắp đặt và bàn giao lại", "Thời hạn báo chấm dứt", "Thời gian khắc phục vi phạm", "Mức phạt/bồi thường và cách tính", "Điều kiện khấu trừ và hoàn trả tiền cọc", "Số ngày hoàn trả tiền cọc", "Điều kiện chậm thanh toán/lãi chậm trả"]:
    field(x)
subsection("1", "Quyền lựa chọn hoặc điều khoản riêng (nếu có)")
for x in ["Loại quyền/điều khoản", "Bên được thực hiện", "Điều kiện và nội dung", "Thời điểm/thời hạn thực hiện", "Diện tích liên quan", "Chi phí hoặc cách tính", "Tài liệu thông báo kèm theo"]:
    field(x)
p("Nếu có nhiều quyền hoặc điều khoản riêng, hai bên ghi tiếp theo các mục trên và ký xác nhận.")
signatures()

annex(5, "ĐẦU MỐI LIÊN HỆ VÀ VĂN BẢN SỬA ĐỔI")
subsection("1", "Đầu mối Bên A")
for x in ["Họ tên", "Vai trò/chức vụ", "Đơn vị", "Điện thoại", "Email", "Địa chỉ nhận thông báo"]:
    field(x)
subsection("2", "Đầu mối Bên B")
for x in ["Họ tên", "Vai trò/chức vụ", "Đơn vị", "Điện thoại", "Email", "Địa chỉ nhận thông báo"]:
    field(x)
subsection("3", "Văn bản sửa đổi, bổ sung sau khi ký (chỉ ghi khi phát sinh)")
for x in ["Số/ký hiệu văn bản", "Ngày đề nghị", "Ngày có hiệu lực", "Bên đề nghị/thực hiện", "Nội dung thay đổi", "Tài liệu kèm theo"]:
    field(x)
p("Văn bản sửa đổi có hiệu lực theo nội dung và chữ ký hợp lệ của hai bên.", before=8)
signatures()

doc.core_properties.title = "Mẫu hợp đồng thuê văn phòng trống"
doc.core_properties.subject = "Mẫu hợp đồng để rà soát dữ liệu OCR"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
